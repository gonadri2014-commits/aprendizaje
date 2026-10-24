"""Genera cotizador.xlsx para el servicio Empresas 2.0.

Uso: python3 generar_cotizador.py
Celdas azules = datos a cargar. El resto se calcula con fórmulas.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

AZUL = Font(color="0000FF")
NEGRITA = Font(bold=True)
TITULO = Font(bold=True, size=14)
INPUT = PatternFill("solid", fgColor="DDEBF7")
TOTAL = PatternFill("solid", fgColor="FFF2CC")
HEADER = PatternFill("solid", fgColor="1F3864")
BLANCO = Font(bold=True, color="FFFFFF")
FINO = Side(style="thin", color="BFBFBF")
BORDE = Border(top=FINO, bottom=FINO, left=FINO, right=FINO)
USD = '"USD" #,##0'

SISTEMAS = [
    ("Excel / sin sistema", 1.0),
    ("Colppy / Xubio / Contabilium / Alegra", 1.0),
    ("Odoo", 1.1),
    ("Tango Gestión", 1.25),
    ("Bejerman / Flexxus / Holistor", 1.3),
    ("SAP Business One", 1.5),
    ("Sistema propio / a medida", 1.7),
]

# (módulo, horas base, incluir por defecto)
MODULOS = [
    ("Conciliación bancaria automática", 30, "SI"),
    ("Dashboard financiero", 45, "SI"),
    ("P&L / EERR automático", 30, "SI"),
    ("Cashflow proyectado", 30, "NO"),
    ("Budget automático", 35, "NO"),
    ("Impuestos (IVA, IIBB, ret./perc.)", 35, "NO"),
    ("Integración con sistema/ERP", 25, "SI"),
]

wb = Workbook()

# --- Parámetros ---
par = wb.active
par.title = "Parametros"
par["A1"] = "Sistema"
par["B1"] = "Factor complejidad"
for c in ("A1", "B1"):
    par[c].font = BLANCO
    par[c].fill = HEADER
for i, (nombre, factor) in enumerate(SISTEMAS, start=2):
    par.cell(i, 1, nombre)
    par.cell(i, 2, factor).font = AZUL
par.column_dimensions["A"].width = 40
par.column_dimensions["B"].width = 20
ult_sis = len(SISTEMAS) + 1

# --- Cotizador ---
ws = wb.create_sheet("Cotizador", 0)
ws["A1"] = "Cotizador — Empresas 2.0"
ws["A1"].font = TITULO
ws["A2"] = "Celdas azules = datos del cliente. Precios en USD."

datos = [
    ("Cliente", "Empresa Ejemplo SRL", None),
    ("Sistema", SISTEMAS[3][0], None),
    ("Bancos / cuentas", 3, None),
    ("Sociedades", 1, None),
    ("Movimientos bancarios por mes", 800, "#,##0"),
    ("Tarifa hora (USD)", 50, USD),
    ("Horas ahorradas al cliente por mes", 60, "#,##0"),
    ("Costo hora del personal del cliente (USD)", 12, USD),
    ("Meses de ahorro a cobrar (3 a 6)", 4, "0"),
    ("Abono mensual (% de implementación)", 0.12, "0%"),
]
for i, (etq, val, fmt) in enumerate(datos, start=4):
    ws.cell(i, 1, etq)
    c = ws.cell(i, 2, val)
    c.font, c.fill, c.border = AZUL, INPUT, BORDE
    if fmt:
        c.number_format = fmt
# Filas: B4 cliente, B5 sistema, B6 bancos, B7 sociedades, B8 movs, B9 tarifa,
# B10 horas ahorradas, B11 costo hora, B12 meses, B13 % abono

dv = DataValidation(type="list", formula1=f"=Parametros!$A$2:$A${ult_sis}")
ws.add_data_validation(dv)
dv.add("B5")

ws["A15"] = "Factor sistema"
ws["B15"] = f"=VLOOKUP(B5,Parametros!$A$2:$B${ult_sis},2,FALSE)"
ws["A16"] = "Factor volumen"
ws["B16"] = "=1+0.1*MAX(B6-1,0)+0.15*MAX(B7-1,0)+IF(B8>2000,0.3,IF(B8>500,0.15,0))"
ws["A17"] = "Factor total"
ws["B17"] = "=B15*B16"
for c in ("B15", "B16", "B17"):
    ws[c].number_format = "0.00"

# Tabla de módulos
fila = 19
for col, h in enumerate(["Módulo", "Incluir", "Horas base", "Horas ajustadas", "Precio por horas"], 1):
    c = ws.cell(fila, col, h)
    c.font, c.fill = BLANCO, HEADER
dv_si = DataValidation(type="list", formula1='"SI,NO"')
ws.add_data_validation(dv_si)
ini = fila + 1
for i, (mod, horas, incl) in enumerate(MODULOS, start=ini):
    ws.cell(i, 1, mod)
    c = ws.cell(i, 2, incl)
    c.font, c.fill = AZUL, INPUT
    dv_si.add(c)
    c = ws.cell(i, 3, horas)
    c.font = AZUL
    ws.cell(i, 4, f'=IF(B{i}="SI",C{i}*$B$17,0)').number_format = "0"
    ws.cell(i, 5, f"=D{i}*$B$9").number_format = USD
    for col in range(1, 6):
        ws.cell(i, col).border = BORDE
fin = ini + len(MODULOS) - 1

r = fin + 2
resumen = [
    ("Total horas estimadas", f"=SUM(D{ini}:D{fin})", "0"),
    ("Precio piso por horas", f"=SUM(E{ini}:E{fin})", USD),
    ("Precio por valor (ahorro × meses)", f"=B10*B11*B12", USD),
    ("PRECIO IMPLEMENTACIÓN (mayor de ambos)", None, USD),
    ("Adelanto 50 %", None, USD),
    ("Saldo a la entrega 50 %", None, USD),
    ("ABONO MENSUAL (mínimo USD 150)", None, USD),
    ("Ahorro anual estimado para el cliente", "=B10*B11*12", USD),
]
for i, (etq, f, fmt) in enumerate(resumen):
    ws.cell(r + i, 1, etq).font = NEGRITA
    ws.cell(r + i, 2, f).number_format = fmt
piso, valor, impl = r + 1, r + 2, r + 3
ws.cell(impl, 2, f"=ROUND(MAX(B{piso},B{valor}),-2)")
ws.cell(impl + 1, 2, f"=B{impl}*0.5")
ws.cell(impl + 2, 2, f"=B{impl}*0.5")
ws.cell(impl + 3, 2, f"=MAX(150,ROUND(B{impl}*B13,-1))")
for fr in (impl, impl + 3):
    for col in (1, 2):
        ws.cell(fr, col).fill = TOTAL

ws.column_dimensions["A"].width = 44
ws.column_dimensions["B"].width = 38
for col in "CDE":
    ws.column_dimensions[col].width = 18
ws["B4"].alignment = Alignment(horizontal="left")

wb.save("cotizador.xlsx")
print("cotizador.xlsx generado")
